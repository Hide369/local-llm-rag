"""Excelブックのテキスト抽出。

1シート=1ユニットにする。PDFのページ・PPTXのスライドと同じく、シートは
書き手が引いた区切りそのものだからである。長いシートは後段の chunk_units が
CHUNK_SIZE で分割する。

表は1本の文字列に潰す時点で行と列の関係を失う。セルを ' | ' で、行を改行で
つなぐのは、見出し行との対応を読み手（および回答を組み立てるLLM）が復元
できるようにするためで、区切りを空白にすると値が地続きになって復元できない。
"""
from pathlib import Path

from openpyxl import load_workbook

from ingest.image_text import describe_image, has_caption, has_ocr
from ingest.models import SHEET, ParsedUnit

_CELL_SEPARATOR = " | "


def _row_text(row) -> str:
    """1行を1本の文字列にする。未入力セルは空文字として詰める。

    openpyxl は未入力セルに None を返す。str() をそのまま通すと本文に 'None'
    が並び、その語が索引されて全シートに共通の無意味な語ができる。
    """
    cells = ["" if value is None else str(value).strip() for value in row]
    while cells and not cells[-1]:
        # 末尾の空セルは列幅の都合で無限に続くことがある。落とさないと
        # 1行が ' |  |  | ...' で埋まり、本文が区切り記号だけになる。
        cells.pop()
    return _CELL_SEPARATOR.join(cells)


def _image_bytes(image) -> bytes | None:
    """openpyxl の画像オブジェクトからバイト列を取り出す。

    実測（openpyxl 3.1.5）では、ファイルから読んだブックの image.ref は
    BytesIO になる。一方、その場で組み立てたブックではパスや PIL の Image に
    なりうるため、両方を受ける。

    _images は openpyxl の私的APIである。公開APIに画像を取り出す口が無い。
    退避先（zipfileで xl/media と xl/drawings のrelsを辿る）は設計書9.2に
    記録してある。
    """
    ref = getattr(image, "ref", None)
    if ref is None:
        return None
    if hasattr(ref, "getvalue"):
        return ref.getvalue()
    if hasattr(ref, "save"):
        import io

        buffer = io.BytesIO()
        ref.save(buffer, format=getattr(ref, "format", None) or "PNG")
        return buffer.getvalue()
    try:
        return Path(str(ref)).read_bytes()
    except OSError:
        return None


def _anchor_key(image):
    """セルの読み順（行→列）に揃えるための整列キー。"""
    anchor = getattr(image, "anchor", None)
    start = getattr(anchor, "_from", None)
    if start is None:
        return (0, 0)
    return (getattr(start, "row", 0), getattr(start, "col", 0))


def _sheet_images(path: Path, caption_image, ocr_bytes) -> dict[str, list[str]]:
    """シート名 → 画像の説明ブロック。

    本文の抽出は read_only=True のままにする（行を逐次読むためメモリを食わない）。
    しかし read_only=True では Worksheet._images が空のままで、画像は一切
    見えない。これが「Excelに貼ったスクリーンショットが読めない」原因だった。
    画像を読むときだけ2回目のロードを行い、読まないときは1回で済ませる。
    """
    book = load_workbook(path)
    try:
        described: dict[str, list[str]] = {}
        for sheet in book.worksheets:
            blocks = []
            for image in sorted(sheet._images, key=_anchor_key):
                blob = _image_bytes(image)
                if blob is None:
                    continue
                text = describe_image(
                    blob,
                    caption_image,
                    ocr_bytes=ocr_bytes,
                    label=f"{path.name} シート「{sheet.title}」",
                )
                if text:
                    blocks.append(text)
            if blocks:
                described[sheet.title] = blocks
        return described
    finally:
        book.close()


def parse_xlsx(path: Path, caption_image=None, on_missing_image=None, ocr_bytes=None) -> list[ParsedUnit]:
    # data_only=True は数式ではなく計算結果を取る。'=SUM(A1:A2)' を索引しても
    # 利用者が読む値はどこにも残らず、検索でも回答でも使えない。
    # Excelが計算結果を保存していないブックでは値が None になるが、その場合に
    # 数式を代わりに入れることはしない。数式は本文ではないためである。
    images = (
        _sheet_images(path, caption_image, ocr_bytes)
        if caption_image is not None or ocr_bytes is not None
        else {}
    )
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        units: list[ParsedUnit] = []
        for sheet in book.worksheets:
            lines = [text for text in (_row_text(row) for row in sheet.iter_rows(values_only=True)) if text]
            blocks = images.get(sheet.title, [])
            if not lines and not blocks:
                # 空シートはブックに残りがちである。ユニットを作ると空の
                # チャンクがDBに入り、どの質問にも弱く一致する。
                # 画像だけのシートは例外である。中身が無いのではなく、
                # 中身が画像の側にある。
                continue
            # 画像は行の後ろに置く。途中へ差し込むと _row_text が作る
            # 「セル | セル」の1行構造が壊れる。
            text = "\n".join([sheet.title, *lines, *blocks])
            units.append(
                ParsedUnit(
                    # シート名を本文の先頭に置く。シート単体で検索に引かれたとき
                    # 何の表なのか分からなくなるのを防ぐためで、md_parser が
                    # 各セクションにH1を付けているのと同じ判断である。
                    text=text,
                    location_type=SHEET,
                    # 通し番号にするのはシート名の重複でチャンクIDが衝突しない
                    # ようにするため。表示にはheadingを使う。
                    location=len(units) + 1,
                    heading=sheet.title,
                    ocr=has_ocr(text),
                    vlm=has_caption(text),
                )
            )
        return units
    finally:
        # read_only=True はファイルハンドルを開いたままにする。閉じないと
        # Windowsではそのブックを別プロセスが開けなくなる。
        book.close()
