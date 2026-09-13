"""Word の雛形。

印は段落単位で扱う。実測 2026-09-13 では、人が編集した docx は1つの段落が
複数の run に割れており、`{{会議名}}` が `['会議名：{{会議', '名}}']` の
2つになっていた。run を1つずつ見て置換すると印が見つからない。

差し込みは、段落のテキストを組み立てて置換し、先頭の run に書き戻して残りの
run を空にする。書式は先頭 run のものになる。印を含まない段落には触らない。
触ると、書式の違う run が先頭のものに潰れる。
"""
import io
from pathlib import Path

import docx

from docgen import marks


def _paragraphs(document):
    """本文・表・ヘッダー・フッターの全段落を返す。

    ヘッダーまで見るのは、会議名をヘッダーに入れる雛形があるためである。
    本文だけ見ていると取りこぼし、しかも「印が無い雛形」と同じ見え方になる。
    """
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _table_paragraphs(table):
    """表のセルの段落。セルの中に表が入ることがあるので再帰する。"""
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def placeholders(path: Path) -> list[str]:
    document = docx.Document(path)
    found: list[str] = []
    for paragraph in _paragraphs(document):
        for name in marks.names(paragraph.text):
            if name not in found:
                found.append(name)
    return found


def fill(path: Path, values: dict[str, str]) -> bytes:
    document = docx.Document(path)
    for paragraph in _paragraphs(document):
        _fill_paragraph(paragraph, values)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fill_paragraph(paragraph, values: dict[str, str]) -> None:
    if not paragraph.runs:
        return
    original = "".join(run.text for run in paragraph.runs)
    replaced = marks.replace(original, values)
    if replaced == original:
        # 印が無い、あるいは値が1つも当たらなかった段落。触らない。
        # 触ると、書式の違う run が先頭 run のものに潰れる。
        return
    # 改行は run.text の setter が <w:br/> に変換する（実測 2026-09-13）。
    paragraph.runs[0].text = replaced
    for run in paragraph.runs[1:]:
        run.text = ""
