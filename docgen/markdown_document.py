"""Markdown を中間構造へ解析し、md / docx / xlsx / pptx へ組む。

雛形なしの生成は LLM に Markdown を1回書かせ、その1つの出力から4形式を作る。
形式ごとにプロンプトを分けると、プロンプトが4本になり、品質のばらつきも4箇所で
別々に面倒を見ることになる。

解析は一度だけ行い、共通の構造（Heading/Paragraph/Bullets/Table/Code/Diagram/
References）にする。形式ごとに Markdown を読み直すと、記法の解釈が4箇所に分かれ、
片方だけ直した状態が例外を出さずに成立する。

ここは docgen/*_template.py と役割が違う。あちらは「他人が作った文書の印を
埋める」役で、書式・レイアウト・ロゴを保つことが仕事である。こちらに原本は
存在しない。混ぜない。
"""
import io
import re
from dataclasses import dataclass, field

import docx
import openpyxl
from docx.shared import Inches, Pt
from pptx import Presentation

from docgen import mermaid


@dataclass
class Heading:
    level: int
    text: str


@dataclass
class Paragraph:
    text: str


@dataclass
class Bullets:
    items: list[str]


@dataclass
class Table:
    header: list[str]
    rows: list[list[str]]


@dataclass
class Code:
    text: str


@dataclass
class Diagram:
    source: str


@dataclass
class References:
    """何を根拠にしたかの記録。コードが組み立て、LLM には書かせない。

    書かせると、渡していないファイルを参照元として並べうる。
    """
    paths: list[str] = field(default_factory=list)
    citations: list[str] = field(default_factory=list)


_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)$")
_FENCE = re.compile(r"^\s*```\s*(\w*)\s*$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in _TABLE_ROW.match(line).group(1).split("|")]


def parse(text: str) -> list:
    """Markdown をブロックの並びにする。

    フェンスの中を先に判定するのは、コード例の中の `#` や `|` を見出しや表と
    取り違えないためである（ingest/parsers/md_parser.py の scan_fences が同じ
    理由でフェンスを先に数えている）。
    """
    blocks: list = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    index = 0
    paragraph: list[str] = []
    bullets: list[str] = []

    def flush():
        nonlocal paragraph, bullets
        if paragraph:
            blocks.append(Paragraph(" ".join(paragraph)))
            paragraph = []
        if bullets:
            blocks.append(Bullets(bullets))
            bullets = []

    while index < len(lines):
        line = lines[index]
        fence = _FENCE.match(line)
        if fence:
            flush()
            language = fence.group(1).lower()
            index += 1
            body: list[str] = []
            while index < len(lines) and not _FENCE.match(lines[index]):
                body.append(lines[index])
                index += 1
            index += 1  # 閉じフェンス
            source = "\n".join(body).strip("\n")
            blocks.append(Diagram(source) if language == "mermaid" else Code(source))
            continue
        if _TABLE_ROW.match(line) and index + 1 < len(lines) and _TABLE_RULE.match(lines[index + 1]):
            flush()
            header = _cells(line)
            index += 2
            rows = []
            while index < len(lines) and _TABLE_ROW.match(lines[index]):
                row = _cells(lines[index])
                # セル数をヘッダーの幅に正規化
                if len(row) < len(header):
                    # 短い行は空文字列でパディング
                    row.extend([""] * (len(header) - len(row)))
                elif len(row) > len(header):
                    # 長い行は余剰セルを削除
                    row = row[:len(header)]
                rows.append(row)
                index += 1
            blocks.append(Table(header, rows))
            continue
        heading = _HEADING.match(line)
        if heading:
            flush()
            blocks.append(Heading(len(heading.group(1)), heading.group(2).strip()))
            index += 1
            continue
        bullet = _BULLET.match(line)
        if bullet:
            if paragraph:
                blocks.append(Paragraph(" ".join(paragraph)))
                paragraph = []
            bullets.append(bullet.group(1).strip())
            index += 1
            continue
        if not line.strip():
            flush()
            index += 1
            continue
        if bullets:
            blocks.append(Bullets(bullets))
            bullets = []
        paragraph.append(line.strip())
        index += 1

    flush()
    return blocks


# 参照したファイルの節の見出し。4形式すべてがこの文字列を使う。
REFERENCES_HEADING = "参照したファイル"


class UnsupportedOutputError(Exception):
    """出力形式として扱えない拡張子を渡された。"""


def _unhandled(block) -> None:
    """描き方の決まっていないブロックが来た。

    ブロックの型は parse() が出す閉じた集合なので、ここへ来るのはプログラムの
    誤りであって利用者の入力ではない。黙って飛ばすと、利用者が受け取る文書から
    本文が消える。例外も警告も出ないため、開いて読むまで誰も気づけない。
    """
    raise UnsupportedOutputError(
        f"描き方の決まっていないブロックです: {type(block).__name__}"
    )


def _table_markdown(block: Table) -> str:
    lines = ["| " + " | ".join(block.header) + " |"]
    lines.append("| " + " | ".join("---" for _ in block.header) + " |")
    lines.extend("| " + " | ".join(row) + " |" for row in block.rows)
    return "\n".join(lines)


# mmdc の既定の PNG は幅800px（約8.3インチ）で、A4縦の段幅（約6.5インチ）より
# 広い。docgen/docx_template.py と同じ値にする。
DIAGRAM_WIDTH = Inches(6.0)


def _add_diagram(document, block, on_diagram_error):
    """図にできれば画像を、できなければ Mermaid のテキストを入れる。

    図にできなかったときに何も入れないと、利用者は成果物を開いても「そこに何か
    あったはず」だと分からない（docgen/mermaid.py の rendered と同じ判断）。
    """
    try:
        image = mermaid.render(block.source)
    except mermaid.MermaidError as error:
        document.add_paragraph(block.source)
        if on_diagram_error is not None:
            # 雛形と違って名前の付いた欄が存在しないので「図」で呼ぶ。
            on_diagram_error("図", str(error))
        return
    document.add_picture(io.BytesIO(image), width=DIAGRAM_WIDTH)


def _build_docx(blocks, on_diagram_error=None) -> tuple[bytes, list[str]]:
    document = docx.Document()
    for block in blocks:
        if isinstance(block, Heading):
            document.add_heading(block.text, level=min(block.level, 9))
        elif isinstance(block, Paragraph):
            document.add_paragraph(block.text)
        elif isinstance(block, Bullets):
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.header))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, block.header):
                cell.text = text
            for row in block.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
        elif isinstance(block, Code):
            paragraph = document.add_paragraph()
            # add_run を明示的に呼ぶことで、空の場合でも run が存在する。
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif isinstance(block, Diagram):
            _add_diagram(document, block, on_diagram_error)
        elif isinstance(block, References):
            document.add_heading(REFERENCES_HEADING, level=2)
            for name in block.paths + block.citations:
                document.add_paragraph(name, style="List Bullet")
        else:
            _unhandled(block)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), []


# python-pptx の既定のレイアウト。0 はタイトル、1 はタイトルと内容。
_TITLE_LAYOUT = 0
_CONTENT_LAYOUT = 1


def _pptx_lines(block) -> list[str]:
    """1つのブロックをスライド本文の行の並びにする。"""
    if isinstance(block, Heading):
        return [block.text]
    if isinstance(block, Paragraph):
        return [block.text]
    if isinstance(block, Bullets):
        return list(block.items)
    if isinstance(block, Table):
        return [" | ".join(block.header)] + [" | ".join(row) for row in block.rows]
    if isinstance(block, Code):
        return block.text.split("\n")
    if isinstance(block, Diagram):
        return block.source.split("\n")
    _unhandled(block)


def _add_slide(presentation, title: str, lines: list[str]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[_CONTENT_LAYOUT])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = lines[0] if lines else ""
    for line in lines[1:]:
        body.add_paragraph().text = line


def _build_pptx(blocks, on_diagram_error=None) -> tuple[bytes, list[str]]:
    """`##` ごとに1スライドにする。

    `###` 以下でスライドを分けないのは、見出しの深さで分けると章立ての書き方
    しだいでスライドが数十枚に膨らむためである。深い見出しは本文の1行にする。

    図は当面 PNG にせず Mermaid のテキストを本文へ入れる。スライドは本文の
    プレースホルダに文字を流し込む作りで、画像を置くと位置と大きさを決める
    判断が要る。on_diagram_error を受け取るのは他の形式と署名を揃えるためである
    （ingest/parsers/__init__.py と同じ方針）。
    """
    presentation = Presentation()
    title = ""
    current: str | None = None
    lines: list[str] = []
    made = False

    def flush():
        nonlocal current, lines, made
        if current is not None or lines:
            _add_slide(presentation, current or title or "", lines)
            made = True
        current, lines = None, []

    for block in blocks:
        if isinstance(block, Heading) and block.level == 1 and not title:
            title = block.text
            slide = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_LAYOUT])
            slide.shapes.title.text = block.text
            made = True
            continue
        if isinstance(block, Heading) and block.level == 2:
            flush()
            current = block.text
            continue
        if isinstance(block, References):
            flush()
            _add_slide(presentation, REFERENCES_HEADING, block.paths + block.citations)
            made = True
            continue
        lines.extend(_pptx_lines(block))
    flush()

    if not made:
        # 見出しも本文も無い回に空のファイルを渡さない。
        _add_slide(presentation, title or "文書", [])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue(), []


# Excel のシート名の制限。31文字以内で、この文字を含められない。
_SHEET_NAME_LIMIT = 31
_SHEET_NAME_FORBIDDEN = re.compile(r"[:\\/?*\[\]]")


def _sheet_name(wanted: str, used: set[str]) -> str:
    """Excel が受け取れるシート名にする。重なったら連番を足す。

    そのまま渡すと保存時に落ちる。落ちると成果物が1つも手に入らない。
    """
    name = _SHEET_NAME_FORBIDDEN.sub("_", wanted).strip() or "表"
    name = name[:_SHEET_NAME_LIMIT]
    if name not in used:
        used.add(name)
        return name
    serial = 2
    while True:
        suffix = f"_{serial}"
        candidate = name[: _SHEET_NAME_LIMIT - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
        serial += 1


def _build_xlsx(blocks, on_diagram_error=None) -> tuple[bytes, list[str]]:
    """表1つにつき1シート。シート名は直前の見出しにする。

    表が1つも無いときは見出しと本文の2列で1シートを出し、警告する。空のブックを
    返さないのは、利用者が受け取るものを空にしないためである。
    """
    book = openpyxl.Workbook()
    book.remove(book.active)
    warnings: list[str] = []
    used: set[str] = set()
    heading = ""
    table_serial = 1
    rows: list[tuple[str, str]] = []

    for block in blocks:
        if isinstance(block, Heading):
            heading = block.text
        elif isinstance(block, Table):
            sheet = book.create_sheet(_sheet_name(heading or f"表{table_serial}", used))
            sheet.append(block.header)
            for row in block.rows:
                sheet.append(row)
            table_serial += 1
            heading = ""
        elif isinstance(block, References):
            sheet = book.create_sheet(_sheet_name(REFERENCES_HEADING, used))
            for name in block.paths + block.citations:
                sheet.append([name])
        elif isinstance(block, Paragraph):
            rows.append((heading, block.text))
            heading = ""
        elif isinstance(block, Bullets):
            for item in block.items:
                rows.append((heading, item))
                heading = ""
        elif isinstance(block, (Code, Diagram)):
            source = block.text if isinstance(block, Code) else block.source
            rows.append((heading, source))
            heading = ""
        else:
            _unhandled(block)

    # 末尾に見出しだけ残っていた場合
    if heading:
        rows.append((heading, ""))

    if not any(isinstance(block, Table) for block in blocks):
        warnings.append(
            "表が1つも書かれなかったため、見出しと本文の2列で出しました。"
            "表が欲しい場合は依頼文で「表で」と指定してください。"
        )
        sheet = book.create_sheet(_sheet_name("本文", used), 0)
        for left, right in rows:
            sheet.append([left, right])

    if not book.sheetnames:
        book.create_sheet(_sheet_name("本文", used))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), warnings


def _build_md(blocks, on_diagram_error=None) -> tuple[bytes, list[str]]:
    parts: list[str] = []
    for block in blocks:
        if isinstance(block, Heading):
            parts.append("#" * block.level + " " + block.text)
        elif isinstance(block, Paragraph):
            parts.append(block.text)
        elif isinstance(block, Bullets):
            parts.append("\n".join(f"- {item}" for item in block.items))
        elif isinstance(block, Table):
            parts.append(_table_markdown(block))
        elif isinstance(block, Code):
            parts.append(f"```\n{block.text}\n```")
        elif isinstance(block, Diagram):
            parts.append(f"```mermaid\n{block.source}\n```")
        elif isinstance(block, References):
            parts.append(f"## {REFERENCES_HEADING}")
            parts.append("\n".join(f"- {name}" for name in block.paths + block.citations))
        else:
            _unhandled(block)
    return ("\n\n".join(parts) + "\n").encode("utf-8"), []


_BUILDERS = {".md": _build_md, ".docx": _build_docx, ".pptx": _build_pptx, ".xlsx": _build_xlsx}

OUTPUT_SUFFIXES = (".md", ".docx", ".xlsx", ".pptx")


def build(blocks, suffix: str, on_diagram_error=None) -> tuple[bytes, list[str]]:
    """ブロックの並びを1つの形式へ組み、(データ, 警告) を返す。

    バイト列を返すのは st.download_button がそれを受け取るためで、中間ファイルを
    作らずに済む（docgen/__init__.py の fill と同じ）。
    """
    builder = _BUILDERS.get(suffix.lower())
    if builder is None:
        raise UnsupportedOutputError(f"出力できない形式です: {suffix}")
    return builder(blocks, on_diagram_error)
