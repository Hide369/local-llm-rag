"""Excel の雛形。

セルの値は文字列1つなので、docx/pptx のような run の分割は起きない。

複数行の値を入れるときだけ折り返しを立てる。常に立てると、雛形が決めた
見た目を勝手に変えることになる。実測 2026-09-13: copy(cell.alignment) して
から wrap_text を立てる形で保存後も保たれる。
"""
import io
from copy import copy
from pathlib import Path

import openpyxl

from docgen import marks


def _cells(book):
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                # 印は文字列にしか現れない。数値セルに置換をかけると型が変わる。
                if isinstance(cell.value, str):
                    yield cell


def placeholders(path: Path) -> list[str]:
    found: list[str] = []
    for cell in _cells(openpyxl.load_workbook(path)):
        for name in marks.names(cell.value):
            if name not in found:
                found.append(name)
    return found


def fill(path: Path, values: dict[str, str]) -> bytes:
    book = openpyxl.load_workbook(path)
    for cell in _cells(book):
        replaced = marks.replace(cell.value, values)
        if replaced == cell.value:
            continue
        cell.value = replaced
        if "\n" in replaced:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
