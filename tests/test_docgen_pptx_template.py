"""PowerPoint の雛形。

実測 2026-09-13: pptx も docx と同じく、段落が複数の run に割れる
（`['会議名：{{会議', '名}}']`）。段落単位で扱う。
"""
import io

import pytest
from pptx import Presentation
from pptx.util import Inches

import docgen


def _blank(tmp_path, name="雛形.pptx"):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide, tmp_path / name


def _textbox(slide, text):
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.paragraphs[0].add_run().text = text
    return box


def _all_text(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def test_a_placeholder_split_across_runs_is_found_and_filled(tmp_path):
    presentation, slide, path = _blank(tmp_path)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = "会議名：{{会議"
    paragraph.add_run().text = "名}}"
    presentation.save(path)

    assert docgen.placeholders(path) == ["会議名"]
    assert _all_text(docgen.fill(path, {"会議名": "第5回"})) == "会議名：第5回"


def test_a_placeholder_in_a_table_cell_is_found_and_filled(tmp_path):
    presentation, slide, path = _blank(tmp_path)
    shape = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(6), Inches(1))
    shape.table.cell(0, 0).text = "出席者"
    shape.table.cell(0, 1).text = "{{出席者}}"
    presentation.save(path)

    assert docgen.placeholders(path) == ["出席者"]
    filled = Presentation(io.BytesIO(docgen.fill(path, {"出席者": "田中"})))
    assert filled.slides[0].shapes[0].table.cell(0, 1).text == "田中"


def test_a_placeholder_inside_a_grouped_shape_is_found(tmp_path):
    """図形はグループにまとめられる。グループの中を見ないと取りこぼす。"""
    presentation, slide, path = _blank(tmp_path)
    first = _textbox(slide, "{{会議名}}")
    second = _textbox(slide, "{{開催日}}")
    slide.shapes.add_group_shape([first, second])
    presentation.save(path)

    assert docgen.placeholders(path) == ["会議名", "開催日"]


def test_an_unfilled_mark_is_left_in_place(tmp_path):
    presentation, slide, path = _blank(tmp_path)
    _textbox(slide, "{{会議名}} / {{決定事項}}")
    presentation.save(path)

    assert _all_text(docgen.fill(path, {"会議名": "第5回"})) == "第5回 / {{決定事項}}"


def test_a_template_without_marks_has_no_placeholders(tmp_path):
    presentation, slide, path = _blank(tmp_path)
    _textbox(slide, "ただの本文")
    presentation.save(path)
    assert docgen.placeholders(path) == []
