"""Excelブックの取り込み。

表は本文と違い、行と列の関係が意味を持つ。埋め込みに渡す時点で1本の文字列に
潰す以上、どの列の値なのかが読み取れる形にしておかないと、数字の羅列になって
検索でも回答でも使えない。
"""
import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image

from ingest.image_text import CAPTION_PREFIX, OCR_PREFIX
from ingest.models import SHEET
from ingest.parsers import SUPPORTED_SUFFIXES, parse
from ingest.parsers.xlsx_parser import parse_xlsx


@pytest.fixture
def book_path(tmp_path):
    book = Workbook()
    first = book.active
    first.title = "商品一覧"
    first.append(["商品コード", "商品名", "価格"])
    first.append(["A-001", "洗濯機", 89800])
    second = book.create_sheet("備考")
    second.append(["納期は三営業日"])
    path = tmp_path / "一覧.xlsx"
    book.save(path)
    return path


@pytest.fixture
def book_with_image_path(tmp_path):
    """シートに画像を1枚貼り付けたブック。

    caption_imageが実際に呼ばれる経路を用意するため book_path とは別に持つ。
    book_path を流用すると「画像が無いので呼ばれない」だけの、恒真になりかねない
    アサーションになってしまう。
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "写真"
    sheet["A1"] = "現地写真"
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (200, 100), "white").save(image_path)
    sheet.add_image(XlsxImage(str(image_path)), "B2")
    path = tmp_path / "写真台帳.xlsx"
    book.save(path)
    return path


def test_each_sheet_becomes_one_unit(book_path):
    units = parse_xlsx(book_path)

    assert len(units) == 2
    assert [unit.location_type for unit in units] == [SHEET, SHEET]
    assert [unit.location for unit in units] == [1, 2]


def test_the_sheet_name_is_carried_in_the_text(book_path):
    """シート単体で引かれたとき、何の表なのかが分からなくなるのを防ぐ。

    Markdownのユニットが見出しにH1を付けているのと同じ理由である。
    """
    assert parse_xlsx(book_path)[0].text.startswith("商品一覧")


def test_a_row_keeps_its_cells_in_one_line(book_path):
    """行の並びが崩れると、どの値がどの列なのか復元できなくなる。"""
    text = parse_xlsx(book_path)[0].text

    assert "商品コード | 商品名 | 価格" in text
    assert "A-001 | 洗濯機 | 89800" in text


def test_the_sheet_name_is_available_for_the_citation(book_path):
    """出典に「一覧.xlsx シート名」と出すため、headingで運ぶ。"""
    assert parse_xlsx(book_path)[0].heading == "商品一覧"


def test_an_empty_sheet_produces_no_unit(tmp_path):
    """空シートはブックに残りがちで、そのままだと空チャンクがDBに入る。"""
    book = Workbook()
    book.active.title = "空"
    book.create_sheet("中身あり").append(["値"])
    path = tmp_path / "空あり.xlsx"
    book.save(path)

    units = parse_xlsx(path)

    assert len(units) == 1
    assert units[0].heading == "中身あり"


def test_empty_cells_do_not_leave_the_string_none(tmp_path):
    """未入力セルは openpyxl が None を返す。str() すると本文に None が並ぶ。"""
    book = Workbook()
    book.active.title = "穴あき"
    book.active.append(["A", None, "C"])
    path = tmp_path / "穴あき.xlsx"
    book.save(path)

    assert "None" not in parse_xlsx(path)[0].text


def test_a_formula_is_stored_as_its_computed_value(tmp_path):
    """数式そのものを索引しても検索の役に立たない。

    data_only を落とすと本文が「=SUM(A1:A2)」になり、利用者が読む値は
    どこにも残らない。Excelが計算結果を保存していない場合は値が無いため、
    その場合にセルが消えることまでを含めてこの挙動である。
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "計算"
    sheet["A1"] = 1
    sheet["A2"] = 2
    sheet["A3"] = "=SUM(A1:A2)"
    path = tmp_path / "計算.xlsx"
    book.save(path)

    assert "=SUM" not in parse_xlsx(path)[0].text


def test_xlsx_is_routed_by_the_registry(book_path):
    """拡張子の登録を忘れると、source/ に置いても黙って無視される。"""
    assert ".xlsx" in SUPPORTED_SUFFIXES
    assert parse(book_path)[0].heading == "商品一覧"


def test_caption_image_is_now_wired_up(book_with_image_path):
    """Task 6で本文へ反映されるようになった（旧: test_caption_image_is_not_yet_wired_up）。

    Task 3時点では read_only=True のせいで ws._images が空になり、画像自体が
    一切見えなかった。Task 6で画像を読むときだけ2回目のロードを行うようにした
    ことで、caption_image が実際に本文へ反映されるようになった。旧テストは
    「反映されない」ことを固定していたが、その前提が崩れたのでここで反転する。
    ocr_bytes を空文字のフェイクにして渡すのは、指定を省くと describe_image が
    ingest.ocr の本物のRapidOCRエンジンを生成してしまうため（既定挙動）。
    """
    text = "\n".join(
        unit.text
        for unit in parse_xlsx(
            book_with_image_path,
            caption_image=lambda _bytes: "説明",
            ocr_bytes=lambda _bytes: "",
        )
    )

    assert "説明" in text


def test_an_embedded_image_is_read(book_with_image_path):
    """read_only=True では画像が一切見えない。これが要望6の原因だった。"""
    text = parse_xlsx(
        book_with_image_path,
        caption_image=lambda _blob: "エラーダイアログです。",
        ocr_bytes=lambda _blob: "コード 0x80070005",
    )[0].text

    assert f"{CAPTION_PREFIX}エラーダイアログです。" in text
    assert f"{OCR_PREFIX}コード 0x80070005" in text


def test_the_image_goes_after_the_rows(book_with_image_path):
    """行の途中へ差し込むと「セル | セル」の1行構造が壊れる。"""
    text = parse_xlsx(
        book_with_image_path,
        caption_image=lambda _blob: "エラーダイアログです。",
        ocr_bytes=lambda _blob: "",
    )[0].text

    assert text.index("現地写真") < text.index("エラーダイアログです。")


def test_the_flags_record_the_engines(book_with_image_path):
    unit = parse_xlsx(
        book_with_image_path,
        caption_image=lambda _blob: "図です。",
        ocr_bytes=lambda _blob: "文字",
    )[0]

    assert unit.vlm is True
    assert unit.ocr is True


def test_a_sheet_with_only_an_image_still_produces_a_unit(tmp_path):
    """空シートを飛ばす判断は「中身が無い」ことが根拠である。画像はその例外になる。"""
    shot = tmp_path / "only.png"
    Image.new("RGB", (300, 180), "blue").save(shot)
    book = Workbook()
    book.active.title = "画像だけ"
    book.active.add_image(XlsxImage(str(shot)), "A1")
    path = tmp_path / "画像だけ.xlsx"
    book.save(path)

    units = parse_xlsx(path, caption_image=lambda _blob: "青い画像です。", ocr_bytes=lambda _b: "")

    assert len(units) == 1
    assert units[0].heading == "画像だけ"
    assert "青い画像です。" in units[0].text


def test_the_workbook_is_opened_once_when_no_image_reader_is_given(book_with_image_path, monkeypatch):
    """画像を読まないときに2回開くと、既存の取り込み時間がただ伸びる。"""
    import ingest.parsers.xlsx_parser as module

    opens = []
    real = module.load_workbook

    def counting(path, **kwargs):
        opens.append(kwargs)
        return real(path, **kwargs)

    monkeypatch.setattr(module, "load_workbook", counting)
    parse_xlsx(book_with_image_path)

    assert len(opens) == 1
    assert opens[0].get("read_only") is True


def test_two_images_on_one_sheet_are_ordered_by_their_anchor_cell(tmp_path):
    """_anchor_key は3段の getattr 連鎖で書かれていて、期待する形が無ければ
    黙って (0, 0) を返す（防御的だが検証は無かった）。2枚の画像を別々の
    セルに置き、行→列の読み順どおりに並ぶかを確かめる。
    """
    import io

    from PIL import Image as PILImage

    def _label(blob: bytes) -> str:
        # どちらの画像かは、再エンコード経路に左右されないようPILで開いて
        # 色そのもので見分ける（バイト列の一致に頼るとエンコード差で崩れる）。
        color = PILImage.open(io.BytesIO(blob)).convert("RGB").getpixel((0, 0))
        return "青い画像です。" if color[2] > color[0] else "赤い画像です。"

    blue_path = tmp_path / "blue.png"
    red_path = tmp_path / "red.png"
    Image.new("RGB", (40, 40), "blue").save(blue_path)
    Image.new("RGB", (40, 40), "red").save(red_path)

    book = Workbook()
    sheet = book.active
    sheet.title = "並び"
    sheet["A1"] = "先頭行"
    # 赤を先に「貼り付ける」が、アンカーは下の行（D10）にする。
    # 挿入順ではなくセル位置で並ぶことを確かめるため、わざと逆にする。
    sheet.add_image(XlsxImage(str(red_path)), "D10")
    sheet.add_image(XlsxImage(str(blue_path)), "A2")
    path = tmp_path / "並び.xlsx"
    book.save(path)

    text = parse_xlsx(path, caption_image=_label, ocr_bytes=lambda _b: "")[0].text

    assert "青い画像です。" in text
    assert "赤い画像です。" in text
    assert text.index("青い画像です。") < text.index("赤い画像です。")


def test_an_empty_sheet_without_images_still_produces_no_unit(tmp_path):
    """画像対応を入れても、本当に空のシートは飛ばし続ける。"""
    book = Workbook()
    book.active.title = "空"
    book.create_sheet("中身あり").append(["値"])
    path = tmp_path / "空あり.xlsx"
    book.save(path)

    units = parse_xlsx(path, caption_image=lambda _blob: "図です。", ocr_bytes=lambda _b: "")

    assert len(units) == 1
    assert units[0].heading == "中身あり"
