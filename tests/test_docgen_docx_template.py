"""Word の雛形。

実測 2026-09-13: 人が編集した docx は1つの段落が複数の run に割れている。
`['会議名：{{会議', '名}}']` のようになるため、run 単位で置換すると印が
見つからない。このファイルのテストはその形を必ず含める。
"""
import io

import docx
import pytest

import docgen


def _save(tmp_path, document, name="雛形.docx"):
    path = tmp_path / name
    document.save(path)
    return path


def _text_of(data: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)


def test_a_placeholder_split_across_runs_is_found(tmp_path):
    """これがこの形式で最も起きやすい壊れ方である。"""
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("会議名：{{会議")
    paragraph.add_run("名}}")
    path = _save(tmp_path, document)
    assert docgen.placeholders(path) == ["会議名"]


def test_a_placeholder_split_across_runs_is_filled(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("会議名：{{会議")
    paragraph.add_run("名}}")
    path = _save(tmp_path, document)
    assert _text_of(docgen.fill(path, {"会議名": "第5回"})) == "会議名：第5回"


def test_a_placeholder_in_a_table_cell_is_found_and_filled(tmp_path):
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "出席者"
    table.cell(0, 1).text = "{{出席者}}"
    path = _save(tmp_path, document)
    assert docgen.placeholders(path) == ["出席者"]
    filled = docx.Document(io.BytesIO(docgen.fill(path, {"出席者": "田中、佐藤"})))
    assert filled.tables[0].cell(0, 1).text == "田中、佐藤"


def test_a_placeholder_in_the_header_is_found_and_filled(tmp_path):
    """会議名をヘッダーに入れる雛形がある。本文だけ見ていると取りこぼす。"""
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "{{会議名}}"
    path = _save(tmp_path, document)
    assert docgen.placeholders(path) == ["会議名"]
    filled = docx.Document(io.BytesIO(docgen.fill(path, {"会議名": "第5回"})))
    assert filled.sections[0].header.paragraphs[0].text == "第5回"


def test_a_multi_line_value_becomes_line_breaks(tmp_path):
    """実測 2026-09-13: run.text に \\n を入れると <w:br/> になる。

    件数が変わる中身は1つの印に複数行で入れる（設計書5節）ので、ここが効く。
    """
    document = docx.Document()
    document.add_paragraph("{{決定事項}}")
    path = _save(tmp_path, document)
    data = docgen.fill(path, {"決定事項": "・A\n・B\n・C"})
    paragraph = docx.Document(io.BytesIO(data)).paragraphs[0]
    assert paragraph.text == "・A\n・B\n・C"
    assert paragraph.runs[0]._element.xml.count("<w:br/>") == 2


def test_an_unfilled_mark_is_left_in_place(tmp_path):
    document = docx.Document()
    document.add_paragraph("会議名：{{会議名}}")
    document.add_paragraph("決定事項：{{決定事項}}")
    path = _save(tmp_path, document)
    assert _text_of(docgen.fill(path, {"会議名": "第5回"})) == (
        "会議名：第5回\n決定事項：{{決定事項}}"
    )


def test_a_paragraph_without_any_mark_is_untouched(tmp_path):
    """印の無い段落に触ると、書式が先頭 run のものに潰れる副作用が出る。"""
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("前半")
    paragraph.add_run("後半").bold = True
    path = _save(tmp_path, document)
    filled = docx.Document(io.BytesIO(docgen.fill(path, {})))
    assert [r.text for r in filled.paragraphs[0].runs] == ["前半", "後半"]


def test_a_template_without_marks_has_no_placeholders(tmp_path):
    document = docx.Document()
    document.add_paragraph("ただの本文")
    assert docgen.placeholders(_save(tmp_path, document)) == []
