"""雛形の保管。

既存のアップロード（scripts/ingest_source.py）は「原本を残さず、DBのチャンク
だけ残す」という設計判断を採っている。雛形はこれと異なる。埋めるために原本
そのものが要るため、ファイルとして残す。
"""
import pytest

import docgen
from docgen import templates as templates_module


def _template(tmp_path, name="議事録.docx"):
    import docx

    path = tmp_path / name
    document = docx.Document()
    document.add_paragraph("{{会議名}}")
    document.save(path)
    return path


def test_a_registered_template_is_listed(tmp_path):
    store = tmp_path / "templates"
    templates_module.register(_template(tmp_path), directory=store)
    assert [p.name for p in templates_module.templates(directory=store)] == ["議事録.docx"]


def test_templates_are_listed_in_name_order(tmp_path):
    store = tmp_path / "templates"
    for name in ("報告書.docx", "議事録.docx", "台帳.xlsx"):
        templates_module.register(_template(tmp_path, name), directory=store)
    listed = [p.name for p in templates_module.templates(directory=store)]
    assert listed == sorted(listed)


def test_registering_the_same_name_replaces_it(tmp_path):
    """利用者にとっては差し替えである。版管理の仕組みは作らない。"""
    store = tmp_path / "templates"
    first = tmp_path / "議事録.docx"
    import docx

    document = docx.Document()
    document.add_paragraph("{{会議名}}")
    document.save(first)
    templates_module.register(first, directory=store)

    document = docx.Document()
    document.add_paragraph("{{会議名}}\n{{決定事項}}")
    document.save(first)
    templates_module.register(first, directory=store)

    assert len(templates_module.templates(directory=store)) == 1
    stored = templates_module.templates(directory=store)[0]
    assert docgen.placeholders(stored) == ["会議名", "決定事項"]


def test_an_unsupported_suffix_is_refused(tmp_path):
    """PDF を登録できてしまうと、選べるのに生成できない雛形が一覧に並ぶ。"""
    store = tmp_path / "templates"
    path = tmp_path / "議事録.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    with pytest.raises(docgen.UnsupportedTemplateError, match="議事録.pdf"):
        templates_module.register(path, directory=store)


def test_removing_a_template_takes_it_off_the_list(tmp_path):
    store = tmp_path / "templates"
    templates_module.register(_template(tmp_path), directory=store)
    templates_module.remove("議事録.docx", directory=store)
    assert templates_module.templates(directory=store) == []


def test_removing_a_name_that_is_not_there_is_not_an_error(tmp_path):
    """2つの画面から同時に消したときに例外を出さない。結果は同じである。"""
    store = tmp_path / "templates"
    templates_module.remove("無い.docx", directory=store)


def test_listing_an_empty_store_returns_nothing(tmp_path):
    """まだ1つも登録していないのは通常の状態であり、例外ではない。"""
    assert templates_module.templates(directory=tmp_path / "templates") == []


def test_a_name_with_a_path_separator_is_refused(tmp_path):
    """名前はファイル名であってパスではない。templates/ の外へ書かせない。"""
    store = tmp_path / "templates"
    with pytest.raises(ValueError, match="ファイル名"):
        templates_module.remove("../vector_store.sqlite3", directory=store)
