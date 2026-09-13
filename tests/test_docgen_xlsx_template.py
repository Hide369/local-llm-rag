"""Excel の雛形。

セルの値は文字列1つなので、docx/pptx のような run の分割は起きない。
代わりに、複数行の値は折り返しを立てないと1行にしか見えない。
"""
import io

import openpyxl

import docgen


def _save(tmp_path, book, name="雛形.xlsx"):
    path = tmp_path / name
    book.save(path)
    return path


def _loaded(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_placeholders_are_found_across_sheets(tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "{{会議名}}"
    second = book.create_sheet("明細")
    second["B2"] = "{{決定事項}}"
    path = _save(tmp_path, book)
    assert docgen.placeholders(path) == ["会議名", "決定事項"]


def test_a_cell_is_filled(tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "会議名：{{会議名}}"
    path = _save(tmp_path, book)
    assert _loaded(docgen.fill(path, {"会議名": "第5回"})).active["A1"].value == "会議名：第5回"


def test_a_multi_line_value_turns_on_wrapping(tmp_path):
    """折り返しを立てないと、改行を入れても画面には1行としか出ない。"""
    book = openpyxl.Workbook()
    book.active["A1"] = "{{決定事項}}"
    path = _save(tmp_path, book)
    filled = _loaded(docgen.fill(path, {"決定事項": "・A\n・B"}))
    assert filled.active["A1"].value == "・A\n・B"
    assert filled.active["A1"].alignment.wrap_text is True


def test_a_single_line_value_does_not_change_the_alignment(tmp_path):
    """折り返しを常に立てると、雛形が決めた見た目を勝手に変えることになる。"""
    book = openpyxl.Workbook()
    book.active["A1"] = "{{会議名}}"
    path = _save(tmp_path, book)
    filled = _loaded(docgen.fill(path, {"会議名": "第5回"}))
    assert not filled.active["A1"].alignment.wrap_text


def test_an_unfilled_mark_is_left_in_place(tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "{{決定事項}}"
    path = _save(tmp_path, book)
    assert _loaded(docgen.fill(path, {})).active["A1"].value == "{{決定事項}}"


def test_a_numeric_cell_is_left_alone(tmp_path):
    """数値セルに文字列置換をかけると型が変わる。印は文字列にしか現れない。"""
    book = openpyxl.Workbook()
    book.active["A1"] = 26
    path = _save(tmp_path, book)
    assert docgen.placeholders(path) == []
    assert _loaded(docgen.fill(path, {})).active["A1"].value == 26
