"""Markdown を中間構造へ解析し、4形式へ組む。

Markdown を一度だけ解析して共通の構造にするのは、形式ごとに読み直すと記法の
解釈が4箇所に分かれるためである。
"""
import io

import docx

from docgen import markdown_document as md
from docgen import mermaid


def _docx_texts(data: bytes) -> list[str]:
    return [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]


def test_headings_keep_their_level():
    blocks = md.parse("# 設計書\n\n## 構成\n")

    assert blocks == [md.Heading(1, "設計書"), md.Heading(2, "構成")]


def test_consecutive_lines_become_one_paragraph():
    """Markdown では空行が段落の区切りである。改行1つで切ると、
    折り返しただけの文が別々の段落になる。"""
    blocks = md.parse("これは1つの\n段落である。\n\n次の段落。\n")

    assert blocks == [md.Paragraph("これは1つの 段落である。"), md.Paragraph("次の段落。")]


def test_bullets_are_collected_into_one_block():
    blocks = md.parse("- 一つ目\n- 二つ目\n")

    assert blocks == [md.Bullets(["一つ目", "二つ目"])]


def test_numbered_lists_are_bullets_too():
    """番号付きも箇条書きとして扱う。docx と pptx で番号を復元する価値より、
    2種類を持ち回る複雑さのほうが大きい。"""
    blocks = md.parse("1. 一つ目\n2. 二つ目\n")

    assert blocks == [md.Bullets(["一つ目", "二つ目"])]


def test_a_table_keeps_its_header_and_rows():
    text = "| 区分 | 日数 |\n| --- | --- |\n| 6か月 | 10日 |\n"

    blocks = md.parse(text)

    assert blocks == [md.Table(["区分", "日数"], [["6か月", "10日"]])]


def test_a_fenced_block_is_code():
    blocks = md.parse("```python\nprint(1)\n```\n")

    assert blocks == [md.Code("print(1)")]


def test_a_mermaid_fence_is_a_diagram():
    """図として描ける可能性があるものを、ただのコードと区別する。"""
    blocks = md.parse("```mermaid\ngraph TD\nA-->B\n```\n")

    assert blocks == [md.Diagram("graph TD\nA-->B")]


def test_a_heading_inside_a_fence_is_not_a_heading():
    """コード例の中の # を見出しにすると、文書の構造が崩れる。"""
    blocks = md.parse("```\n# これはコメント\n```\n")

    assert blocks == [md.Code("# これはコメント")]


def test_an_empty_document_has_no_blocks():
    assert md.parse("   \n\n") == []


def test_a_blank_line_is_what_separates_two_tables():
    """2つの表を分ける方法は空行である。GFM の標準動作として、空行がなければ
    2つのテーブル構文が1つにマージされる。これは仕様であり、バグではない。"""
    # 空行なし：マージされて1つのテーブルになる
    text_no_blank = "| a | b |\n| - | - |\n| 1 | 2 |\n| c | d |\n| - | - |\n| 3 | 4 |\n"
    blocks_no_blank = md.parse(text_no_blank)
    assert blocks_no_blank == [
        md.Table(["a", "b"], [["1", "2"], ["c", "d"], ["-", "-"], ["3", "4"]])
    ]

    # 空行あり：2つの別々のテーブルになる
    text_with_blank = "| a | b |\n| - | - |\n| 1 | 2 |\n\n| c | d |\n| - | - |\n| 3 | 4 |\n"
    blocks_with_blank = md.parse(text_with_blank)
    assert blocks_with_blank == [
        md.Table(["a", "b"], [["1", "2"]]),
        md.Table(["c", "d"], [["3", "4"]])
    ]


def test_a_short_row_is_padded_to_header_width():
    """セル数がヘッダーより少ない行は、空文字列でパディングされる。
    docx/xlsx/pptx レンダラーは固定列数で構築されるため、不揃いな行が
    列ずれを引き起こさないように正規化が必須である。"""
    text = "| a | b |\n| - | - |\n| 1 |\n"

    blocks = md.parse(text)

    assert blocks == [md.Table(["a", "b"], [["1", ""]])]


def test_a_long_row_surplus_cells_are_dropped():
    """セル数がヘッダーより多い行の余剰セルは削除される。
    GFM の仕様に準じ、下流レンダラーが誤配置しない形に整形する。"""
    text = "| a | b |\n| - | - |\n| 1 | 2 | 3 |\n"

    blocks = md.parse(text)

    assert blocks == [md.Table(["a", "b"], [["1", "2"]])]


def test_a_table_with_legitimate_dash_data_rows_keeps_all_rows():
    """データ行に | - | - | が含まれる場合、先読みで表を分けようとすると
    区切り行と見分けが付かず、この行が消える。複数の表でマージを許容する以上、
    シンプルな停止条件（空行まで続ける）が唯一の正解である。"""
    text = "| a | b |\n| - | - |\n| 1 | 2 |\n| - | - |\n| 3 | 4 |\n"

    blocks = md.parse(text)

    assert blocks == [md.Table(["a", "b"], [["1", "2"], ["-", "-"], ["3", "4"]])]


def test_build_md_round_trips_the_blocks():
    blocks = [md.Heading(1, "設計書"), md.Paragraph("本文"), md.Bullets(["一つ目"])]

    data, warnings = md.build(blocks, ".md")

    assert warnings == []
    assert data.decode("utf-8") == "# 設計書\n\n本文\n\n- 一つ目\n"


def test_build_md_keeps_mermaid_as_a_fence():
    """.md は GitHub・VS Code・画面のプレビューが図として表示する。
    PNG にする理由が無い（docgen/md_template.py と同じ判断）。"""
    data, _ = md.build([md.Diagram("graph TD\nA-->B")], ".md")

    assert data.decode("utf-8") == "```mermaid\ngraph TD\nA-->B\n```\n"


def test_build_md_writes_the_references_section():
    blocks = [md.Paragraph("本文"), md.References(["main.go"], ["議事録.docx p.1"])]

    text = md.build(blocks, ".md")[0].decode("utf-8")

    assert "## 参照したファイル" in text
    assert "- main.go" in text
    assert "- 議事録.docx p.1" in text


def test_build_rejects_an_unsupported_suffix():
    import pytest
    with pytest.raises(md.UnsupportedOutputError):
        md.build([md.Paragraph("本文")], ".pdf")


def test_build_raises_on_unhandled_block_type():
    """黙って飛ばすと利用者が受け取る文書から本文が消え、例外も警告も出ないため気づけない。"""
    import pytest
    from dataclasses import dataclass

    @dataclass
    class UnknownBlock:
        content: str

    with pytest.raises(md.UnsupportedOutputError):
        md.build([UnknownBlock("本文")], ".md")


def test_build_docx_writes_headings_and_paragraphs():
    blocks = [md.Heading(1, "設計書"), md.Paragraph("本文")]

    data, warnings = md.build(blocks, ".docx")

    assert warnings == []
    assert "設計書" in _docx_texts(data)
    assert "本文" in _docx_texts(data)


def test_build_docx_writes_a_table():
    blocks = [md.Table(["区分", "日数"], [["6か月", "10日"]])]

    document = docx.Document(io.BytesIO(md.build(blocks, ".docx")[0]))

    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "区分"
    assert document.tables[0].cell(1, 1).text == "10日"


def test_build_docx_writes_the_references_section():
    blocks = [md.References(["main.go"], ["議事録.docx p.1"])]

    texts = _docx_texts(md.build(blocks, ".docx")[0])

    assert md.REFERENCES_HEADING in texts
    assert "main.go" in texts
    assert "議事録.docx p.1" in texts


def test_build_docx_keeps_the_mermaid_text_when_it_cannot_be_drawn(monkeypatch):
    """図にできなかったことを黙って捨てると、利用者は成果物を開くまで
    気づけない。テキストは残し、呼び出し元へ知らせる。"""
    def _fail(source):
        raise mermaid.MermaidError("mmdc を起動できませんでした")

    monkeypatch.setattr(mermaid, "render", _fail)
    seen = []

    data, warnings = md.build(
        [md.Diagram("graph TD\nA-->B")], ".docx",
        lambda name, reason: seen.append((name, reason)),
    )

    assert "graph TD" in "\n".join(_docx_texts(data))
    assert seen and "mmdc" in seen[0][1]


def test_build_docx_raises_on_unhandled_block_type():
    """黙って飛ばすと利用者が受け取る文書から本文が消え、例外も警告も出ないため気づけない。"""
    import pytest
    from dataclasses import dataclass

    @dataclass
    class UnknownBlock:
        content: str

    with pytest.raises(md.UnsupportedOutputError):
        md.build([UnknownBlock("本文")], ".docx")


def test_build_docx_writes_a_code_block():
    """コードブロックのテキストが文書に含まれることを確認する。"""
    blocks = [md.Code("print(1)")]

    texts = _docx_texts(md.build(blocks, ".docx")[0])

    assert "print(1)" in texts


def test_build_docx_handles_empty_code_block():
    """空のフェンスはモデルが書きうる。`runs[0]` を無条件に触ると、そこで生成ごと落ちる。"""
    blocks = [md.Code("")]

    data, warnings = md.build(blocks, ".docx")

    assert warnings == []
    assert data is not None  # 文書が正常に生成されることを確認
