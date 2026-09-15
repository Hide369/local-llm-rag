"""Cowork タブ（雛形からの文書生成）。

実機のOllamaにもネットワークにも触れない。AppTest は rag_chat_app.py を同じ
プロセスで実行するため、本番のストアと本番の templates/ を開かせない。
"""
import json
from pathlib import Path
from unittest.mock import patch

import docx
import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

import ingest.retrieval as retrieval
from docgen import filling
from docgen import markdown_document
from docgen import mermaid
from docgen import source_code
from docgen import templates as templates_module
from ingest import chat
from ingest import reranker as reranker_module
from ingest import store as store_module
from ingest import vlm as vlm_module
from ingest.models import ParsedUnit
from ingest.vector_store import open_store as open_real_store

APP_PATH = Path(__file__).resolve().parent.parent / "rag_chat_app.py"


@pytest.fixture
def app():
    st.cache_resource.clear()
    return AppTest.from_file(str(APP_PATH), default_timeout=60)


@pytest.fixture(autouse=True)
def _no_network():
    with (
        patch.object(reranker_module, "check_reranker", lambda: None),
        patch.object(reranker_module, "rerank", lambda query, texts: [9.0] * len(texts)),
        patch.object(vlm_module, "check_vlm", lambda *a, **k: None),
    ):
        yield


def _notes(app):
    """補足のメッセージ。

    以前は st.warning の黄色枠で並べていたが、生成のたびに画面の大半が警告で
    埋まるため、畳んだ expander の中へ移した。中身は st.write なので markdown
    として出る。
    """
    return list(app.markdown)


def _reads(*paths):
    """read_files を1周だけ呼び、2周目では呼ばないモデル。

    project.gather のループを画面越しに動かすための最小の模擬である。
    """
    rounds = {"count": 0}

    def ask_tools(model, messages, tools, session=None, num_ctx=None):
        rounds["count"] += 1
        if rounds["count"] == 1:
            return {
                "role": "assistant",
                "tool_calls": [
                    {
                        "function": {
                            "name": "read_files",
                            "arguments": {"paths": list(paths)},
                        }
                    }
                ],
            }
        return {"role": "assistant", "content": "読めました"}

    return ask_tools


def _reads_nothing(model, messages, tools, session=None, num_ctx=None):
    """道具を1度も呼ばないモデル。ツリーだけが渡る経路を通す。"""
    return {"role": "assistant", "content": "読みません"}


def _stub_store():
    def factory(*args, **kwargs):
        collection = open_real_store(":memory:")
        collection.add(
            ids=["chunk-1"],
            documents=["第5回 AI活用検討会を開催した。"],
            embeddings=[[0.1, 0.2]],
            metadatas=[{"source": "議事録.docx", "location_type": "section", "location": 1}],
        )
        return collection

    return factory


def _register(directory, name="議事録.docx"):
    directory.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph("{{会議名}}")
    document.add_paragraph("{{決定事項}}")
    document.save(directory / name)


def test_the_mode_toggle_offers_both_modes(app, tmp_path):
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
    ):
        app.run()
    assert not app.exception
    assert "Cowork" in str(app)


def test_cowork_without_any_template_does_not_tell_the_user_to_register_one(app, tmp_path):
    """雛形が0件でも、雛形なしの選択肢があるため行き止まりにならない。

    以前は0件のとき「雛形が登録されていません」という警告だけを出し、生成の
    手段が無い行き止まりだった。雛形なしが正規の選択肢になった今、その警告は
    出ない（雛形なしで生成できることは test_cowork_offers_generating_without_a_template
    が確かめる）。
    """
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()

    assert not app.exception
    assert not any("雛形が登録されていません" in note.value for note in _notes(app))


def test_cowork_offers_the_register_button_before_any_template_exists(app, tmp_path):
    """雛形が0件のときこそ登録ボタンが要る。

    ボタンを「雛形が1つ以上あるとき」の側に置いていたため、初めて Cowork を
    開いた人の画面には雛形を登録する手段が1つも無く、警告文だけが存在しない
    ボタンを指していた。登録できなければ Cowork は何もできない（実機で確認
    2026-09-13）。
    """
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        button = next(b for b in app.button if b.label == "雛形を登録・削除")
        button.click().run()

    assert not app.exception
    # 押した先が登録の画面であることまで見る。ボタンだけ出て何も開かなければ
    # 行き止まりは解けていない。
    assert app.file_uploader(key="template_files") is not None


def test_cowork_without_a_template_does_not_silently_answer_as_chat(app, tmp_path):
    """雛形を選ばないまま Cowork で送っても、黙ってチャットの回答を返さない。

    元の分岐は「Cowork かつ雛形あり」以外を全部チャットの else 節に落として
    いたため、雛形0件で送信すると通常のチャット回答が返っていた。利用者は
    Cowork のつもりで読むため、どこから来た答えなのかを取り違える。雛形なしが
    正規の選択肢になった今も、この経路がチャットの else 節へ落ちないことに
    変わりはない。

    messages が空であることだけでは、リクエストが黙って握りつぶされた場合
    （何も生成されない）とも区別が付かない。雛形なしの生成が実際に結果を
    残したこと（ダウンロードボタンが出て、エラーが無いこと）まで確かめる。
    """
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "# 議事録\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert app.session_state["messages"] == []
    assert app.download_button
    assert not app.error


def test_cowork_shows_the_error_when_generation_fails(app, tmp_path):
    """PromptTooLongError等の失敗は、埋まらない欄を残すだけでなく画面に伝える。

    黙って落とすと、利用者は生成が止まったこと自体に気づけない
    （_generate_document は結果を st.session_state.cowork_result に積むだけで
    直接 st.error を呼ばないため、例外経路がその積み方から漏れていないかを
    別途押さえる）。
    """
    store = tmp_path / "templates"
    _register(store)
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(
            filling,
            "fill_values",
            side_effect=filling.PromptTooLongError("添付と検索結果が長すぎます"),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert any("添付と検索結果が長すぎます" in error.value for error in app.error)


def test_cowork_generates_a_file_and_offers_it_for_download(app, tmp_path):
    store = tmp_path / "templates"
    _register(store)
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {"会議名": "第5回 AI活用検討会", "決定事項": "・継続"}, ensure_ascii=False
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert app.download_button


def test_cowork_lists_the_marks_it_could_not_fill(app, tmp_path):
    """埋まらなかった欄を出さないと、利用者は雛形を開くまで気づけない。"""
    store = tmp_path / "templates"
    _register(store)
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {"会議名": "第5回 AI活用検討会"}, ensure_ascii=False
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert any("決定事項" in note.value for note in _notes(app))


def test_cowork_uses_a_larger_context_size(app, tmp_path):
    """8192 のままだと、文字起こしを添付した瞬間に文脈からこぼれる。"""
    store = tmp_path / "templates"
    _register(store)
    seen = []

    def ask_json(model, prompt, session=None, num_ctx=None):
        seen.append(num_ctx)
        return json.dumps({"会議名": "第5回"}, ensure_ascii=False)

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_json", ask_json),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("議事録を作って").run()

    from docgen import filling

    assert seen == [filling.GENERATION_NUM_CTX]


def test_cowork_does_not_search_the_documentation_unless_it_is_asked(app, tmp_path):
    """既定で技術ドキュメントまで引くと、1回の生成ごとに英訳のLLM呼び出しが
    1回と検索が1本、要らないまま増える（30〜60秒）。"""
    store = tmp_path / "templates"
    _register(store)
    opened = []
    collection = _stub_store()

    def factory(db_path, *args, **kwargs):
        opened.append(str(db_path))
        return collection(db_path)

    with (
        patch.object(store_module, "open_store", factory),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {"会議名": "第5回"}, ensure_ascii=False
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert not any("docs_store" in path for path in opened)


def test_cowork_tells_the_user_when_a_diagram_could_not_be_drawn(app, tmp_path):
    """黙って Mermaid のテキストが入っていると、利用者は成果物を開くまで
    気づけない。"""
    store = tmp_path / "templates"
    _register(store)
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(
            mermaid, "render", side_effect=mermaid.MermaidError("mmdc が見つかりません")
        ),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {
                    "会議名": "```mermaid\nsequenceDiagram\n  利用者->>画面: 生成\n```",
                    "決定事項": "・継続",
                },
                ensure_ascii=False,
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("設計書を作って").run()

    assert not app.exception
    assert any("図にできなかった" in note.value for note in _notes(app))


def test_cowork_reports_a_broken_template_instead_of_crashing(app, tmp_path):
    """壊れた雛形（.docx の拡張子で中身が違うファイル）を Cowork で送っても、
    生のトレースバックで止まらない。templates.register は拡張子しか見ないため、
    これは異常系ではなく起こりうる誤操作である。

    generating が戻ることも合わせて確かめる。戻らないと、以後どの操作をしても
    同じ分岐に入って同じ例外で落ち、利用者はブラウザーのセッションを捨てる
    以外の出口を失う。
    """
    store = tmp_path / "templates"
    store.mkdir(parents=True, exist_ok=True)
    (store / "壊れた.docx").write_bytes(b"not a docx file")
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "壊れた.docx").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert any("開けませんでした" in error.value for error in app.error)
    assert app.session_state["generating"] is False
    assert app.session_state["pending_question"] is None


def test_template_dialog_shows_a_broken_template_but_still_offers_delete(app, tmp_path):
    """壊れた雛形を消せる唯一の画面がここである。一覧の描画で落ちると、
    削除ボタンごと出せなくなり利用者は雛形を消せなくなる。
    """
    store = tmp_path / "templates"
    store.mkdir(parents=True, exist_ok=True)
    (store / "壊れた.docx").write_bytes(b"not a docx file")
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        button = next(b for b in app.button if b.label == "雛形を登録・削除")
        button.click().run()

    assert not app.exception
    assert any("開けません" in markdown.value for markdown in app.markdown)
    assert app.button(key="remove_壊れた.docx") is not None


def test_cowork_does_not_leave_the_request_in_the_chat_history(app, tmp_path):
    """Cowork の依頼文をチャット履歴に残さない。

    残っていると、次のチャットの contextual_query（ingest/retrieval.py）が
    返信の無い直前の user メッセージとして Cowork の依頼文を拾い、検索クエリが
    「第5回会議の議事録を作って <新しい質問>」のように汚れる。Cowork の依頼は
    結果パネル（render_cowork_result）が受け持つので、チャット履歴に残す必要は
    ない。
    """
    store = tmp_path / "templates"
    _register(store)
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {"会議名": "第5回"}, ensure_ascii=False
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert app.session_state["messages"] == []


def test_cowork_attachments_are_captioned_through_the_vlm(app, tmp_path):
    """既存のアップロード経路（upload_dialog）と同じ判断を添付にも適用する。

    caption_image を渡さずに parse を呼ぶと、画面の図やスクリーンショットを
    添付しても説明が付かず、OCR で拾えた文字だけになる。
    """
    store = tmp_path / "templates"
    _register(store)
    seen = []

    def fake_parse(path, caption_image=None, on_missing_image=None):
        seen.append(caption_image)
        return [ParsedUnit(text="本文", location_type="document", location=1)]

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch("ingest.parsers.parse", fake_parse),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {"会議名": "第5回"}, ensure_ascii=False
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.file_uploader(key="cowork_files").set_value(
            ("画面.png", b"\x89PNG\r\n\x1a\n", "image/png")
        )
        app.run()
        app.chat_input[0].set_value("議事録を作って").run()

    assert not app.exception
    assert seen == [vlm_module.caption_image]


NO_TEMPLATE = "（雛形なし）"


def test_cowork_offers_generating_without_a_template(app, tmp_path):
    """雛形を作る手間のほうが大きい仕事がある。雛形が0件でも画面は成立する。"""
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()

    assert not app.exception
    assert NO_TEMPLATE in app.selectbox(key="template").options


def test_cowork_offers_source_code_as_an_output_format(app, tmp_path):
    """コードを書かせたい回に、文書の形式しか選べないと出口が無い。"""
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()

    assert not app.exception
    options = app.selectbox(key="output_suffix").options
    assert ".go" in options
    assert ".cs" in options
    assert ".md" in options


def test_cowork_generates_go_source_without_the_markdown_path(app, tmp_path):
    """.go はコードであって文書ではない。

    Markdown の解析器へ通すと `//` コメントや `|` を含む行が見出しや表として
    解釈される。ここでは経路が分かれていること（source_code 側が呼ばれ、
    フェンスが剥がれ、参照コメントが先頭に付くこと）を確かめる。
    """
    built = {}
    real_build = source_code.build

    def spy_build(code, paths, citations, suffix):
        data, warnings = real_build(code, paths, citations, suffix)
        built["data"] = data
        return data, warnings

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "```go\npackage main\n```"),
        patch.object(chat, "ask_json", lambda *a, **k: "{}"),
        patch.object(source_code, "build", spy_build),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.selectbox(key="output_suffix").set_value(".go").run()
        app.chat_input[0].set_value("Go で取り込み処理を書いて").run()

    assert not app.exception
    text = built["data"].decode("utf-8")
    assert "```" not in text
    assert text.rstrip().endswith("package main")
    assert app.download_button


def test_writing_source_into_the_project_warns_about_the_build(app, tmp_path):
    """generated_docs/ は走査からは外してあるが、言語のツールチェーンは別の規則で動く。

    文書と違い、ソースコードは置いた場所がビルドの対象になる。`go build ./...`
    は generated_docs/ の .go も拾うため、package 宣言が食い違えばその時点で
    ビルドが壊れる。黙って置くと、原因がこの機能だと気づくのが難しい。
    """
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "main.go").write_text("package main\n", encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "package main\n"),
        patch.object(chat, "ask_tools", _reads("main.go")),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.selectbox(key="output_suffix").set_value(".go").run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("Go で書いて").run()

    assert not app.exception
    written = list((folder / "generated_docs").glob("*.go"))
    assert len(written) == 1
    assert any("ビルド" in note.value for note in _notes(app))


def test_cowork_without_a_template_generates_markdown(app, tmp_path):
    """雛形なしの生成が実際に .md のバイト列を組み立てることまで確かめる。

    AppTest の DownloadButton は proto.url しか持たず、st.download_button に
    渡した bytes を属性としては公開していない（実データは AppTest が
    実行のたびに作り直す MediaFileManager が握っており、app.run() の外から
    覗けない）。そこで markdown_document.build に実装をそのまま素通りさせる
    スパイを挟み、実際に生成へ渡されたバイト列を横取りする。
    """
    built = {}
    real_build = markdown_document.build

    def spy_build(blocks, suffix, on_diagram_error=None):
        data, warnings = real_build(blocks, suffix, on_diagram_error)
        built["data"] = data
        return data, warnings

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
        patch.object(chat, "ask_json", lambda *a, **k: "{}"),
        patch.object(markdown_document, "build", spy_build),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.selectbox(key="output_suffix").set_value(".md").run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert app.download_button
    assert built["data"].decode("utf-8").startswith("# 設計書")


def test_cowork_without_a_template_stops_when_there_is_no_evidence(app, tmp_path):
    """根拠が1つも無ければ、呼んでも中身の無い文書が出るだけである。"""
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.checkbox(key="cowork_internal").set_value(False).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert any("根拠" in error.value for error in app.error)


def test_cowork_warns_when_an_attachment_yields_no_text(app, tmp_path):
    """説明文もOCR文字も得られなかった添付を黙って通さない。

    空の本文がプロンプトに載っても利用者には何も伝わらないため、画面に
    警告として出す。
    """
    store = tmp_path / "templates"
    _register(store)

    def fake_parse(path, caption_image=None, on_missing_image=None):
        return []

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch("ingest.parsers.parse", fake_parse),
        patch.object(
            chat,
            "ask_json",
            lambda model, prompt, session=None, num_ctx=None: json.dumps(
                {"会議名": "第5回"}, ensure_ascii=False
            ),
        ),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.file_uploader(key="cowork_files").set_value(
            ("画面.png", b"\x89PNG\r\n\x1a\n", "image/png")
        )
        app.run()
        app.chat_input[0].set_value("議事録を作って").run()

    assert not app.exception
    assert any("本文を取り出せませんでした" in note.value for note in _notes(app))


def test_cowork_reads_the_project_folder_and_writes_the_result_into_it(app, tmp_path):
    """成果物はダウンロードだけでなくフォルダにも残す。"""
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "main.go").write_text("package main\n", encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_tools", _reads("main.go")),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    written = list((folder / "generated_docs").glob("*.md"))
    assert len(written) == 1
    assert "# 設計書" in written[0].read_text(encoding="utf-8")


def test_the_generated_document_lists_the_files_it_used(app, tmp_path):
    """何を根拠にしたかを後から辿れるようにする。"""
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "main.go").write_text("package main\n", encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_tools", _reads("main.go")),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    written = list((folder / "generated_docs").glob("*.md"))
    assert len(written) == 1
    text = written[0].read_text(encoding="utf-8")
    assert "## 参照したファイル" in text
    assert "main.go" in text


def _embed_query_must_not_be_called(*args, **kwargs):
    raise AssertionError("embed_query was called before the folder was validated")


def test_a_folder_that_does_not_exist_stops_before_calling_the_model(app, tmp_path):
    """呼んでから落ちると、利用者は30〜60秒待たされたうえで何も受け取れない。

    パスの打ち間違いを伝えるのに、埋め込み検索と翻訳の LLM を通す理由がない。
    社内資料を参照するチェックは既定でオンのままにしておき、
    `retrieval.embed_query` が「呼ばれたら失敗する」関数に差し替えてある
    ことで、フォルダの検証（`docgen_project.tree`）が検索より先に実行される
    順序をロックする。以前はこの検証が検索より後にあり、フォルダを打ち
    間違えただけでも埋め込み検索（と技術ドキュメント参照時は翻訳のLLM
    呼び出しまで）を1回無駄に払ってから初めてエラーが出ていた。
    """
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", _embed_query_must_not_be_called),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(tmp_path / "無い")).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert any("フォルダ" in error.value for error in app.error)


def test_a_folder_with_no_supported_file_names_the_formats_it_accepts(app, tmp_path):
    """「取り込める形式のファイルがありません」だけでは原因が分からない。

    実測 2026-09-15: `.py` が対象から漏れていたため Python のプロジェクトを
    指定した回がこの文言で止まったが、利用者には「なぜ0件なのか」を知る手段が
    画面上に無く、拡張子の漏れだと分かるまで時間がかかった。受け付ける形式を
    文言に含めれば、同じことが次に起きたときはその場で分かる。

    この文言は警告であって、エラーではない。0件でも生成は続く
    （test_an_empty_project_folder_warns_but_still_generates を参照）。
    """
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "index.ts").write_text("export const x = 1;\n", encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    message = "\n".join(note.value for note in _notes(app))
    assert ".py" in message
    assert ".md" in message


def test_a_failed_selection_still_lets_the_template_path_see_the_listing(app, tmp_path):
    """選択が0件でも、雛形ありの経路もツリーだけは載せる。

    以前は _collect_project_files が「一覧だけ渡します」と警告する一方で、
    _generate_document 側は tree_text をアンダースコア変数で受けて捨てて
    いたため、警告の内容と実際の挙動が食い違っていた（設計書5節の
    フォールバックは freeform 経路にしか効いていなかった）。
    """
    store = tmp_path / "templates"
    _register(store)
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "main.go").write_text("package main\n", encoding="utf-8")

    seen_attachments = []
    real_fill_values = filling.fill_values

    def spy_fill_values(placeholders, question, sources, attachments, ask):
        seen_attachments.append(attachments)
        return real_fill_values(placeholders, question, sources, attachments, ask)

    def ask_json(model, prompt, session=None, num_ctx=None):
        return json.dumps({"会議名": "第5回"}, ensure_ascii=False)

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", store),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        # 道具を1度も呼ばないモデル。読めたファイルは0件になり、ツリーだけが残る。
        patch.object(chat, "ask_tools", _reads_nothing),
        patch.object(chat, "ask_json", ask_json),
        patch.object(filling, "fill_values", spy_fill_values),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(store / "議事録.docx").run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("第5回会議の議事録を作って").run()

    assert not app.exception
    assert seen_attachments
    names = [name for name, _ in seen_attachments[0]]
    assert "（プロジェクトのファイル一覧）" in names
    listing = next(text for name, text in seen_attachments[0] if name == "（プロジェクトのファイル一覧）")
    assert "main.go" in listing


def test_a_project_folder_too_big_for_the_listing_warns_with_the_exact_count(app, tmp_path):
    """一覧がTREE_BUDGET_CHARSを超えるフォルダでは、モデルに見えていない
    ファイルがあることを画面で伝える。伝えないと、選ばれなかったファイルが
    「関係が無いから選ばれなかった」のか「一覧に載らなかったから選べな
    かった」のか利用者には区別が付かない。
    """
    from docgen import project as project_module

    folder = tmp_path / "myproject"
    folder.mkdir()
    entries_count = 400
    for i in range(entries_count):
        (folder / f"f{i:04d}.go").write_text("package main\n", encoding="utf-8")
    entries = project_module.tree(folder)
    _listing, expected_omitted = project_module.tree_text(entries)
    assert expected_omitted > 0  # このテストの前提（実測で確かめる）

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(chat, "ask_tools", _reads_nothing),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.checkbox(key="cowork_internal").set_value(False).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert any(str(expected_omitted) in note.value for note in _notes(app))


def test_a_project_folder_alone_is_enough_evidence(app, tmp_path):
    """フォルダだけを指定した回は正当な使い方であり、止めてはならない。"""
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "main.go").write_text("package main\n", encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(chat, "ask_tools", _reads("main.go")),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.checkbox(key="cowork_internal").set_value(False).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert not app.error


def test_an_empty_project_folder_warns_but_still_generates(app, tmp_path):
    """空のフォルダを指定しても止めない。

    プロジェクトフォルダは資料源であると同時に**出力先**でもある。これから
    作るプロジェクトを指定して書かせたい回に、中身が無いという理由で何も
    受け取れないのは筋が通らない。以前はここで生成ごと止めていた。
    """
    folder = tmp_path / "empty_project"
    folder.mkdir()

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert not app.error
    assert any("取り込める形式のファイル" in w.value for w in _notes(app))
    # 出力先としては使える。成果物はこのフォルダに残る。
    assert list((folder / "generated_docs").glob("*.md"))


def test_a_folder_that_is_not_a_directory_still_stops(app, tmp_path):
    """存在しないパスは打ち間違いである。書き出しもできないので進めない。"""
    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", _embed_query_must_not_be_called),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(tmp_path / "無い")).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    assert any("フォルダ" in error.value for error in app.error)


def test_when_the_model_chooses_nothing_the_files_are_read_anyway(app, tmp_path):
    """モデルが道具を呼ばなくても、根拠を0件にしない。

    実測 2026-09-15: gpt-oss:20b は read_files を1度も呼ばず、成果物は毎回
    ファイル一覧だけを根拠に書かれていた。警告は正しかったが、正しいだけで
    利用者は何もできなかった。
    """
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "main.go").write_text("package main\n", encoding="utf-8")
    # 全部が予算に収まると、モデルへ聞かずに全部読む経路へ流れてしまう。
    # 往復が起きる大きさにしたうえで、モデルが何も選ばない状況を作る。
    (folder / "かさ増し.md").write_text("あ" * 20_000, encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_tools", _reads_nothing),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    written = next((folder / "generated_docs").glob("*.md"))
    # 読んだ以上、参照したファイルに名前が出る。
    assert "main.go" in written.read_text(encoding="utf-8")
    # 「選べませんでした」で終わる補足は出さない。何を読んだかまで言う。
    for note in _notes(app):
        assert "ファイル一覧だけを渡します" not in note.value


def test_the_supplementary_messages_are_not_yellow_boxes(app, tmp_path):
    """生成のたびに黄色い枠が積み上がると、成果物が画面の下へ押し流される。

    実測 2026-09-15: ingest/ を指定した回は「読まなかったファイル」だけで
    28件・452字の枠になっていた。捨てはせず、畳んだ場所へ移す。
    """
    folder = tmp_path / "myproject"
    folder.mkdir()
    (folder / "index.ts").write_text("export const x = 1;\n", encoding="utf-8")

    with (
        patch.object(store_module, "open_store", _stub_store()),
        patch.object(templates_module, "TEMPLATE_DIR", tmp_path / "templates"),
        patch.object(retrieval, "embed_query", lambda *a, **k: [0.1, 0.2]),
        patch.object(chat, "ask_text", lambda *a, **k: "# 設計書\n\n本文\n"),
    ):
        app.run()
        app.segmented_control[0].set_value("Cowork").run()
        app.selectbox(key="template").set_value(NO_TEMPLATE).run()
        app.text_input(key="project_folder").set_value(str(folder)).run()
        app.chat_input[0].set_value("設計書を書いて").run()

    assert not app.exception
    # 黄色い枠は1つも出さない。
    assert not app.warning
    # 中身は畳んだ場所に残る。捨てているわけではない。
    assert any("取り込める形式のファイル" in note.value for note in _notes(app))
